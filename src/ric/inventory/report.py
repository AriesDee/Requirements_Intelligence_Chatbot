"""Excel export for the standalone Requirements Inventory Explorer."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ric.models import Requirement

from .analyzer import InventoryStats

_HDR_FILL = PatternFill("solid", fgColor="1A4A8A")
_HDR_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
_GAP_FILL = PatternFill("solid", fgColor="FFF3CD")  # amber for gap rows
_ALT_FILL = PatternFill("solid", fgColor="F0F4F8")


def _hdr(ws, col: int, row: int, value: str) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.fill = _HDR_FILL
    c.font = _HDR_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_inventory_excel(
    requirements: list[Requirement],
    stats: InventoryStats,
    path: Path,
) -> None:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Full Inventory ────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Full Inventory"

    headers = [
        "ID", "Source", "Paygroup", "Section", "Category",
        "Scope Type", "Labor Agreements", "Hour Types",
        "Earn Codes", "Adj Earn Codes", "Description",
    ]
    for col, h in enumerate(headers, 1):
        _hdr(ws1, col, 1, h)

    for row_idx, r in enumerate(requirements, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else None
        vals = [
            r.id, r.source, r.paygroup, r.section, r.category or "",
            r.scope_type,
            ", ".join(r.labor_agreements),
            ", ".join(r.hour_types),
            ", ".join(r.earn_codes),
            ", ".join(r.adj_earn_codes),
            r.description,
        ]
        for col, v in enumerate(vals, 1):
            c = ws1.cell(row=row_idx, column=col, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=(col == len(headers)))
            if fill:
                c.fill = fill

    col_widths = [28, 8, 12, 28, 20, 12, 32, 22, 22, 22, 70]
    for col, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(col)].width = w
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = ws1.dimensions
    ws1.row_dimensions[1].height = 22

    # ── Sheet 2: Quality Gaps ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Quality Gaps")
    gap_headers = ["ID", "Source", "Paygroup", "Section", "Issue"]
    for col, h in enumerate(gap_headers, 1):
        _hdr(ws2, col, 1, h)

    for row_idx, g in enumerate(stats.gaps, 2):
        for col, key in enumerate(gap_headers, 1):
            c = ws2.cell(row=row_idx, column=col, value=g[key])
            c.fill = _GAP_FILL
            c.alignment = Alignment(vertical="top")

    for col, w in enumerate([28, 8, 12, 28, 50], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    # ── Sheet 3: Summary ───────────────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")
    _hdr(ws3, 1, 1, "Metric")
    _hdr(ws3, 2, 1, "Value")

    summary_rows = [
        ("Total Requirements", stats.total),
        ("Quality Gaps", len(stats.gaps)),
        ("Sections", len(stats.by_section)),
        ("Labor Agreements Referenced", len(stats.by_la)),
        ("", ""),
        ("By Source", ""),
    ] + [(f"  {k}", v) for k, v in stats.by_source.items()] + [
        ("", ""),
        ("By Scope Type", ""),
    ] + [(f"  {k}", v) for k, v in stats.by_scope_type.items()] + [
        ("", ""),
        ("Top Sections", "Count"),
    ] + list(stats.by_section.items())[:15] + [
        ("", ""),
        ("Top Labor Agreements", "Count"),
    ] + list(stats.by_la.items())[:15]

    for row_idx, (label, value) in enumerate(summary_rows, 2):
        ws3.cell(row=row_idx, column=1, value=label)
        ws3.cell(row=row_idx, column=2, value=value)

    ws3.column_dimensions["A"].width = 35
    ws3.column_dimensions["B"].width = 15

    try:
        wb.save(path)
    finally:
        wb.close()
