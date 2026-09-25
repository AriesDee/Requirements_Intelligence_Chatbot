from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ric.ioc.models import IOCParseResult
from ric.matcher.models import ChangeMatch

# Palette
_C_HEADER = "1A3A5C"   # dark navy — header rows
_C_DIRECT = "FFF3E0"   # pale orange — direct match rows
_C_INDIRECT = "E8F0FE" # pale blue — indirect match rows
_C_NO_MATCH = "F5F5F5" # light grey — unmatched rows
_C_SECTION = "E8EEF5"  # blue-grey — section label rows
_C_STATUS = "FFFDE7"   # pale yellow — status column

_FONT_BASE = Font(name="Calibri", size=10)
_FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
_FONT_TITLE = Font(name="Calibri", size=14, bold=True, color=_C_HEADER)
_FONT_SECTION = Font(name="Calibri", size=10, bold=True, color=_C_HEADER)

_ALIGN_WRAP = Alignment(wrap_text=True, vertical="top")
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _set_col_widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Summary sheet
# ---------------------------------------------------------------------------

def _write_summary(ws, ioc: IOCParseResult, results: list[ChangeMatch]) -> None:
    ws.title = "IOC Summary"
    ws.sheet_view.showGridLines = False

    def kv(label: str, value: str, row: int) -> None:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _FONT_SECTION
        lc.alignment = _ALIGN_WRAP
        vc = ws.cell(row=row, column=2, value=value)
        vc.font = _FONT_BASE
        vc.alignment = _ALIGN_WRAP

    def section(title: str, row: int) -> None:
        c = ws.cell(row=row, column=1, value=title)
        c.font = _FONT_TITLE
        ws.cell(row=row, column=2)

    r = 1
    ws.cell(r, 1, "Requirements Intelligence Chatbot — Impact Analysis Report").font = _FONT_TITLE
    r += 2

    section("IOC Details", r); r += 1
    kv("Contract IDs", ", ".join(ioc.contract_ids), r); r += 1
    kv("Labor Agreement IDs", ", ".join(ioc.labor_agreement_ids), r); r += 1
    kv("Effective Date", ioc.effective_date or "—", r); r += 1
    kv("Union Locals", ", ".join(ioc.union_locals) or "—", r); r += 1
    r += 1

    total_changes = len(results)
    tk_changes = sum(1 for cm in results if cm.change.timekeeping_relevant)
    total_matches = sum(len(cm.matched_requirements) for cm in results)
    direct = sum(
        1 for cm in results for m in cm.matched_requirements if m.relevance == "direct"
    )
    indirect = total_matches - direct
    unmatched = sum(1 for cm in results if not cm.matched_requirements)

    section("Analysis Summary", r); r += 1
    kv("Total IOC Changes", str(total_changes), r); r += 1
    kv("Timekeeping-Relevant Changes", str(tk_changes), r); r += 1
    kv("Requirements Matched", str(total_matches), r); r += 1
    kv("  — Direct Matches", str(direct), r); r += 1
    kv("  — Indirect Matches", str(indirect), r); r += 1
    kv("Unmatched Changes", str(unmatched), r); r += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60


# ---------------------------------------------------------------------------
# Impact analysis sheet
# ---------------------------------------------------------------------------

_HEADERS = [
    "#", "Change Type", "TK Relevant", "Eff. Date", "LA Codes",
    "Change Summary", "Req ID", "Src", "Section",
    "Req Description", "Relevance", "Rationale", "Delta Description", "Status",
]

_COL_WIDTHS = [5, 18, 10, 12, 14, 40, 28, 6, 22, 45, 12, 50, 55, 15]


def _header_fill():
    return _fill(_C_HEADER)


def _write_impact(ws, results: list[ChangeMatch], req_lookup: dict | None = None) -> None:
    ws.title = "Impact Analysis"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    # Header row
    for col, h in enumerate(_HEADERS, 1):
        c = ws.cell(1, col, h)
        c.font = _FONT_HEADER
        c.fill = _header_fill()
        c.alignment = _ALIGN_CENTER

    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}1"

    row = 2
    for idx, cm in enumerate(results, 1):
        change = cm.change

        base_vals = [
            idx,
            change.change_type.value,
            "Y" if change.timekeeping_relevant else "N",
            change.effective_date or "",
            ", ".join(change.labor_agreements),
            change.summary,
        ]

        if cm.matched_requirements:
            first = True
            for m in cm.matched_requirements:
                fill_color = _C_DIRECT if m.relevance == "direct" else _C_INDIRECT
                row_fill = _fill(fill_color)

                req = (req_lookup or {}).get(m.requirement_id)
                vals = base_vals + [
                    m.requirement_id,
                    f"{req.source}-{req.paygroup}" if req else "",
                    req.section if req else "",
                    req.description if req else "",
                    m.relevance,
                    m.rationale,
                    m.delta_description or "",
                    "",  # Status
                ]
                for col, v in enumerate(vals, 1):
                    c = ws.cell(row, col, v)
                    c.font = _FONT_BASE
                    c.fill = row_fill
                    c.alignment = _ALIGN_WRAP
                    if col == len(vals):  # Status column
                        c.fill = _fill(_C_STATUS)

                # Draw a top border on the first match row of each change group
                if first:
                    from openpyxl.styles import Border, Side
                    thin = Side(style="thin", color="AAAAAA")
                    for col in range(1, len(_HEADERS) + 1):
                        ws.cell(row, col).border = Border(top=thin)
                    first = False

                row += 1
        else:
            row_fill = _fill(_C_NO_MATCH)
            vals = base_vals + ["", "", "", "", "no match",
                                cm.no_match_reason or "", "", ""]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row, col, v)
                c.font = Font(name="Calibri", size=10, color="888888")
                c.fill = row_fill
                c.alignment = _ALIGN_WRAP
            row += 1

    _set_col_widths(ws, _COL_WIDTHS)

    # Row heights for data rows
    for r in range(2, row):
        ws.row_dimensions[r].height = 42


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

_SEV_COLORS = {"high": "FFCDD2", "medium": "FFF9C4", "low": "E8F5E9"}
_SEV_FONT_COLORS = {"high": "B71C1C", "medium": "F57F17", "low": "1B5E20"}

_FLAG_HEADERS = ["#", "Change Type", "TK", "Change Summary", "Severity", "Flag Type", "Description", "Recommendation"]
_FLAG_COL_WIDTHS = [5, 18, 6, 40, 10, 22, 50, 55]


def _write_flags(ws, results: list[ChangeMatch]) -> None:
    ws.title = "Risk Flags"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    for col, h in enumerate(_FLAG_HEADERS, 1):
        c = ws.cell(1, col, h)
        c.font = _FONT_HEADER
        c.fill = _fill(_C_HEADER)
        c.alignment = _ALIGN_CENTER

    ws.auto_filter.ref = f"A1:{get_column_letter(len(_FLAG_HEADERS))}1"

    row = 2
    for idx, cm in enumerate(results, 1):
        ch = cm.change
        flags = getattr(cm, "risk_flags", [])
        if not flags:
            continue
        for flag in flags:
            sev = flag.severity.value if hasattr(flag.severity, "value") else str(flag.severity)
            bg = _SEV_COLORS.get(sev, "FFFFFF")
            fc = _SEV_FONT_COLORS.get(sev, "000000")
            row_fill = _fill(bg)
            vals = [
                idx,
                ch.change_type.value,
                "Y" if ch.timekeeping_relevant else "N",
                ch.summary,
                sev.upper(),
                flag.flag_type,
                flag.description,
                flag.recommendation,
            ]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row, col, v)
                c.font = Font(name="Calibri", size=10,
                              bold=(col == 5),
                              color=fc if col == 5 else "000000")
                c.fill = row_fill
                c.alignment = _ALIGN_WRAP
            row += 1

    _set_col_widths(ws, _FLAG_COL_WIDTHS)
    for r in range(2, row):
        ws.row_dimensions[r].height = 52


def write_excel(
    ioc: IOCParseResult,
    results: list[ChangeMatch],
    path: Path,
    requirements=None,
) -> None:
    """Write the full impact analysis workbook to path."""
    req_lookup = {r.id: r for r in requirements} if requirements else {}
    wb = openpyxl.Workbook()
    _write_summary(wb.active, ioc, results)
    ws2 = wb.create_sheet()
    _write_impact(ws2, results, req_lookup)
    ws3 = wb.create_sheet()
    _write_flags(ws3, results)
    wb.save(path)
